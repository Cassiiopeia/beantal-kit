"""문서 읽기: pptx/docx/xlsx/pdf 를 공통 구조(dict)로 만든다.

- 평문 파일은 python 라이브러리로 읽는다.
- DRM(사내 문서 보안) 파일은 Windows 에서 PowerPoint·Word·Excel 을 통해 읽기 전용으로 열어 메모리에서만 구조를 뽑는다.
  복호화된 사본을 디스크에 만들지 않는다.
"""
import json
import re
import os
import platform
import re
import subprocess
from pathlib import Path

EMU_PER_CM = 360000


def cm(v):
    return round(v / EMU_PER_CM, 2) if v is not None else None


def clean(s):
    return re.sub(r"\s+", " ", s or "").strip()


# ---------------------------------------------------------------- pptx (python)

def _shape_record(sh, MSO):
    rec = {
        "name": sh.name,
        "left": cm(sh.left), "top": cm(sh.top), "width": cm(sh.width), "height": cm(sh.height),
        "kind": "other", "text": "", "paras": [],
    }
    t = sh.shape_type
    if t == MSO.PICTURE:
        rec["kind"] = "picture"
    elif getattr(sh, "has_chart", False) and sh.has_chart:
        rec["kind"] = "chart"
    elif getattr(sh, "has_table", False) and sh.has_table:
        rec["kind"] = "table"
        rows = []
        for r in sh.table.rows:
            rows.append([clean(c.text) for c in r.cells])
        rec["table"] = rows
        rec["text"] = " | ".join(" | ".join(r) for r in rows)
    elif t == MSO.LINE or "Connector" in type(sh).__name__:
        rec["kind"] = "line"
    elif t == MSO.AUTO_SHAPE:
        rec["kind"] = "autoshape"
    elif t == MSO.TEXT_BOX:
        rec["kind"] = "textbox"
    elif t == MSO.PLACEHOLDER:
        rec["kind"] = "placeholder"
    if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
        paras = []
        for p in sh.text_frame.paragraphs:
            txt = clean(p.text)
            if not txt:
                continue
            size = next((r.font.size.pt for r in p.runs if r.font.size), None)
            font = next((r.font.name for r in p.runs if r.font.name), None)
            bold = any(bool(r.font.bold) for r in p.runs)
            color = None
            for r in p.runs:
                try:
                    if r.font.color and r.font.color.type is not None and r.font.color.rgb is not None:
                        color = str(r.font.color.rgb)
                        break
                except Exception:
                    pass
            paras.append({"text": txt, "size": size, "font": font, "bold": bold, "color": color, "level": p.level})
        rec["paras"] = paras
        if not rec["text"]:
            rec["text"] = " ".join(x["text"] for x in paras)
    return rec


def _walk(shapes, MSO):
    for sh in shapes:
        if sh.shape_type == MSO.GROUP:
            yield from _walk(sh.shapes, MSO)
        else:
            yield sh


def read_pptx_python(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE as MSO
    prs = Presentation(path)
    fonts_theme = set()
    slides = []
    for i, s in enumerate(prs.slides, 1):
        shapes = [_shape_record(sh, MSO) for sh in _walk(s.shapes, MSO)]
        notes = ""
        if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
            notes = clean(s.notes_slide.notes_text_frame.text)
        slides.append({"n": i, "layout": s.slide_layout.name, "notes": notes, "shapes": shapes})
    return {
        "format": "pptx", "via": "python",
        "size_cm": [cm(prs.slide_width), cm(prs.slide_height)],
        "masters": len(prs.slide_masters),
        "slides": slides,
    }


# ---------------------------------------------------------------- pptx (PowerPoint, DRM)

def read_pptx_powerpoint(path):
    data = run_app_script("drm_read_pptx.ps1", ["-Path", str(path)], "powerpoint", timeout=600)
    if data.get("error"):
        return data
    data["format"] = "pptx"
    data["via"] = "powerpoint"
    return data


def read_pptx(path):
    kind = file_kind(path)
    if kind == "missing":
        return {"error": "missing", "message": f"파일이 없습니다: {path}"}
    if kind == "drm":
        d = read_pptx_powerpoint(path)
    else:
        d = read_pptx_python(path)
    d["drm"] = kind == "drm"
    d["path"] = str(path)
    d["mb"] = round(Path(path).stat().st_size / 1e6, 1)
    return d


# ================================================================ 아무 파일이나 읽기 (read_any)
#
# 확장자 표를 늘리는 대신 "어떻게 생긴 파일인가"(앞부분 바이트·zip 안 이름)와 "무엇이 열 수 있나"(설치된 앱)로
# 시도 순서를 만들고, 하나가 실패하면 다음 방법으로 넘어간다. 전부 실패하면 시도 기록(tried)과
# 다음에 해 볼 일(hints)을 돌려줘서 에이전트가 스스로 다른 길을 고르게 한다.
#
# 새 형식을 붙일 때: 아래 READERS 에 함수 하나와 guess_routes() 에 한 줄을 더한다.

TEXT_EXT = {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".html", ".htm", ".log", ".yml", ".yaml", ".ini"}
WORD_EXT = {".docx", ".doc", ".docm", ".dotx", ".dot", ".rtf", ".odt"}
EXCEL_EXT = {".xlsx", ".xls", ".xlsm", ".xlsb", ".xltx", ".ods"}
PPT_EXT = {".pptx", ".ppt", ".pptm", ".potx", ".ppsx", ".odp"}
HWP_EXT = {".hwp", ".hwpx"}
APP_PROC = {"word": "WINWORD.EXE", "excel": "EXCEL.EXE", "powerpoint": "POWERPNT.EXE", "hangul": "Hwp.exe"}
APP_NAME = {"word": "Word", "excel": "Excel", "powerpoint": "PowerPoint", "hangul": "한글"}


def sniff(path):
    """파일 생김새: drm · zip-docx · zip-xlsx · zip-pptx · zip-hwpx · zip · pdf · ole · text · unknown."""
    with open(path, "rb") as f:
        head = f.read(8)
    if head.startswith(b"BMS"):
        return "drm"
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "ole"  # 옛 Office(.doc·.xls·.ppt) · .hwp
    if head.startswith(b"PK"):
        import zipfile
        try:
            names = zipfile.ZipFile(path).namelist()
        except Exception:
            return "zip"
        for prefix, k in (("word/", "zip-docx"), ("xl/", "zip-xlsx"), ("ppt/", "zip-pptx"), ("Contents/", "zip-hwpx")):
            if any(n.startswith(prefix) for n in names):
                return k
        return "zip"
    try:
        with open(path, "rb") as f:
            f.read(4096).decode("utf-8")
        return "text"
    except UnicodeDecodeError:
        return "unknown"


def file_kind(path):
    """'plain' | 'drm' | 'missing' | 'unknown' — 기존 호출부 호환용."""
    if not Path(path).exists():
        return "missing"
    s = sniff(path)
    return "drm" if s == "drm" else ("unknown" if s == "unknown" else "plain")


def guess_routes(ext, look):
    """시도할 읽기 방법 순서. 평문이면 python 먼저(빠르고 토큰 적음), 안 되면 설치된 앱으로."""
    routes = []
    native = {"zip-docx": "docx", "zip-xlsx": "xlsx", "zip-pptx": "pptx", "zip-hwpx": "hwpx", "pdf": "pdf", "text": "text"}
    if look in native:
        routes.append(native[look])
    # 확장자로 짐작한 앱 (DRM·옛 형식·라이브러리 실패 모두 여기로)
    if ext in WORD_EXT or look == "zip-docx":
        routes.append("word")
    if ext in EXCEL_EXT or look == "zip-xlsx" or ext in {".csv", ".tsv"}:
        routes.append("excel")
    if ext in PPT_EXT or look == "zip-pptx":
        routes.append("powerpoint")
    if ext in HWP_EXT or look == "zip-hwpx":
        routes.append("hangul")
    if ext == ".pdf" or look == "pdf":
        routes.append("word")  # Word 가 pdf 를 문단으로 바꿔 연다 (DRM pdf 는 막힐 수 있음)
    if ext in TEXT_EXT and "text" not in routes:
        routes.append("text")
    if ext in HWP_EXT:
        return [r for r in dict.fromkeys(routes)]  # 한글 문서는 한글 앱으로만 (다른 앱은 글자가 깨짐)
    # 확장자를 모를 때: 생김새로 한 번 더 짐작하고, 그래도 없으면 앱을 차례로
    if not routes or look in ("drm", "ole", "unknown"):
        for r in ("word", "excel", "powerpoint", "hangul", "text"):
            if r not in routes:
                routes.append(r)
    seen = []
    return [r for r in routes if not (r in seen or seen.append(r))]


def _app_pids(image):
    try:
        out = subprocess.run(["tasklist", "/FI", f"IMAGENAME eq {image}", "/FO", "CSV", "/NH"], capture_output=True, timeout=30)
        rows = out.stdout.decode("utf-8", errors="replace").splitlines()
        return {int(r.split('","')[1]) for r in rows if r.startswith('"') and '","' in r}
    except Exception:
        return set()


def run_app_script(script_name, args, app, timeout=180):
    """앱 COM 을 쓰는 ps1 을 실행해 JSON 을 받는다. 시간이 넘으면 이번에 새로 뜬 앱 프로세스만 정리한다
    (사용자가 열어 둔 창은 건드리지 않는다)."""
    if platform.system() != "Windows":
        return {"error": "needs_windows", "message": f"{APP_NAME.get(app, app)} 로 읽는 방법은 Windows 에서만 됩니다."}
    image = APP_PROC.get(app)
    before = _app_pids(image) if image else set()
    script = Path(__file__).resolve().parent.parent / script_name
    cmd = ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script), *args]
    try:
        out = subprocess.run(cmd, capture_output=True, timeout=timeout)
    except subprocess.TimeoutExpired:
        if image:
            for pid in _app_pids(image) - before:
                subprocess.run(["taskkill", "/PID", str(pid), "/F"], capture_output=True)
        return {"error": "app_timeout", "message": f"{APP_NAME.get(app, app)} 가 {timeout}초 안에 끝나지 않았습니다 (숨은 대화상자·보안 프로그램 차단 가능)."}
    raw = out.stdout.decode("utf-8", errors="replace").strip()
    try:
        data = json.loads(raw[raw.index("{"):]) if "{" in raw else None
    except ValueError:
        data = None
    if not data:
        err = out.stderr.decode("utf-8", errors="replace").strip().splitlines()
        return {"error": "app_failed", "message": f"{APP_NAME.get(app, app)} 로 읽지 못했습니다: {' '.join(err[:2])[:300] or raw[-300:]}"}
    return data


def _rows(rows):
    """PowerShell JSON 은 원소 하나짜리 배열을 문자열로 펴기도 한다 — 표를 항상 2차원 목록으로 맞춘다."""
    return [[str(c) for c in r] if isinstance(r, list) else [str(r)] for r in rows or []]


# ---------------------------------------------------------------- 읽기 방법들 (python)

def read_docx(path):
    import docx
    d = docx.Document(path)
    paras = []
    for p in d.paragraphs:
        t = clean(p.text)
        if t:
            paras.append({"text": t, "style": p.style.name})
    tables = []
    for tb in d.tables:
        tables.append([[clean(c.text) for c in r.cells] for r in tb.rows])
    return {"format": "docx", "paras": paras, "tables": tables}


def read_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else clean(str(v)) for v in row]
            while vals and not vals[-1]:
                vals.pop()
            if any(vals):
                rows.append(vals)
        sheets[ws.title] = rows
    return {"format": "xlsx", "sheets": sheets}


def _pdf_visual(pg, text):
    """글자로 안 읽히는 그림·도형이 있는 쪽인지 본다. (이미지 개수, 벡터 도형 연산 수)로 판단한다."""
    try:
        n_img = len(pg.images)
    except Exception:
        n_img = 0
    n_vec = 0
    try:
        data = pg.get_contents().get_data() if pg.get_contents() else b""
        n_vec = len(re.findall(rb"(?m)\s(?:re|c|l)\s*$", data))
    except Exception:
        pass
    reasons = []
    if n_img:
        reasons.append(f"이미지 {n_img}개")
    if n_vec >= 40:
        reasons.append(f"도형·선 {n_vec}개 (표·흐름도일 수 있음)")
    if len(text) < 80:
        reasons.append("글자가 거의 없음")
    return {"images": n_img, "vectors": n_vec, "reasons": reasons}


def read_pdf(path):
    from pypdf import PdfReader
    r = PdfReader(path)
    pages = []
    for i, pg in enumerate(r.pages, 1):
        text = (pg.extract_text() or "").strip()
        v = _pdf_visual(pg, text)
        pages.append({"n": i, "text": text, "images": v["images"], "visual": bool(v["images"] or v["vectors"] >= 40 or len(text) < 80),
                      "visual_reasons": v["reasons"]})
    if not any(p["text"] for p in pages):
        raise ValueError("글자가 없는 pdf (스캔 이미지일 수 있음)")
    return {"format": "pdf", "pages": pages, "visual_pages": [p["n"] for p in pages if p["visual"]]}


def read_text(path):
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "cp949", "utf-16"):
        try:
            txt = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("글자 인코딩을 알 수 없습니다")
    return {"format": "pdf", "pages": [{"n": 1, "text": txt}]}  # 줄 목록 출력은 pdf 와 같은 모양으로 쓴다


def read_hwpx(path):
    """hwpx = zip 안의 XML. 문단(<hp:p>)마다 글자(<hp:t>)를 모은다."""
    import zipfile
    z = zipfile.ZipFile(path)
    secs = sorted(n for n in z.namelist() if re.match(r"Contents/section\d+\.xml", n))
    lines = []
    for n in secs:
        xml = z.read(n).decode("utf-8", errors="replace")
        for para in re.findall(r"<hp:p[ >].*?</hp:p>", xml, flags=re.S):
            t = clean(" ".join(re.findall(r"<hp:t[^>]*>(.*?)</hp:t>", para, flags=re.S)))
            if t:
                lines.append(re.sub(r"<[^>]+>", "", t))
    return {"format": "pdf", "pages": [{"n": 1, "text": "\n".join(lines)}]}


# ---------------------------------------------------------------- 읽기 방법들 (설치된 앱, DRM 포함)

def _office(path, kind, app, timeout=180):
    data = run_app_script("drm_read_office.ps1", ["-Path", str(path), "-Kind", kind], app, timeout)
    if data.get("error"):
        return data
    if kind == "docx":
        data["format"] = "docx"
        data["paras"] = [p if isinstance(p, dict) else {"text": str(p), "style": ""} for p in data.get("paras") or []]
        data["tables"] = [_rows(t) for t in data.get("tables") or []]
    elif kind == "xlsx":
        data["format"] = "xlsx"
        data["sheets"] = {name: _rows(rows) for name, rows in (data.get("sheets") or {}).items()}
    else:
        data["format"] = "pdf"
        data["pages"] = data.get("pages") or []
    return data


def _via_word(path):
    return _office(path, "pdf" if path.suffix.lower() == ".pdf" else "docx", "word", 120 if path.suffix.lower() == ".pdf" else 300)


READERS = {
    # 이름: (함수, 평문 python 인가)
    "pptx": (lambda p: read_pptx_python(p), True),
    "docx": (read_docx, True),
    "xlsx": (read_xlsx, True),
    "pdf": (read_pdf, True),
    "hwpx": (read_hwpx, True),
    "text": (read_text, True),
    "word": (_via_word, False),
    "excel": (lambda p: _office(p, "xlsx", "excel", 300), False),
    "powerpoint": (lambda p: read_pptx_powerpoint(p), False),
    "hangul": (lambda p: _office(p, "hwp", "hangul", 180), False),
}


def _hints(p, look, tried):
    h = []
    if look == "drm":
        h.append("DRM 파일입니다. 이 파일을 여는 프로그램이 이 PC 에 있으면 `read --via <word|excel|powerpoint|hangul>` 로 그 프로그램을 지정해 다시 시도하세요.")
    if p.suffix.lower() in HWP_EXT:
        h.append("한글(HWP)은 한컴오피스가 설치된 PC 에서만 읽힙니다. 없으면 사용자에게 한글에서 열어 필요한 부분을 복사해 붙여 달라고 하거나 PDF·DOCX 로 저장해 달라고 하세요.")
    if p.suffix.lower() == ".pdf" or look == "pdf":
        h.append("pdf 가 안 열리면: 스캔 이미지 pdf 라면 `snapshot` 처럼 화면으로 보고 읽거나, 사용자에게 Acrobat 에서 전체 선택·복사해 붙여 달라고 하세요.")
    if any(t.get("error") == "app_timeout" for t in tried):
        h.append("앱이 시간 안에 안 끝났습니다. 사용자에게 그 앱 창에 대화상자(보안·변환 안내)가 떠 있는지 확인해 달라고 하세요.")
    h.append("그래도 안 되면 이 파일을 어떤 프로그램으로 여는지 사용자에게 묻고, 필요한 부분을 대화로 붙여 받으세요. 복호화 사본을 디스크에 만들지는 않습니다.")
    return h


def _garbled(d):
    """글자가 깨져 나온 결과(엉뚱한 앱이 바이너리를 억지로 연 경우)를 가려낸다."""
    texts = []
    if d.get("pages"):
        texts = [p.get("text", "") for p in d["pages"]]
    elif d.get("paras"):
        texts = [p.get("text", "") for p in d["paras"]]
    s = "".join(texts)[:5000]
    if len(s) < 20:
        return False
    # 한글·영문·숫자·공백·흔한 기호 비율이 낮고 물음표·대체문자가 많으면 깨진 것
    good = sum(1 for ch in s if ch.isalnum() or ch.isspace() or ch in ".,:;()[]-_/%·•※▣①②③④⑤→<>|+*&'\"!~#@=")
    bad = sum(1 for ch in s if ch in "?�\x00")
    return bad / len(s) > 0.15 or good / len(s) < 0.6


def read_any(path, via=None):
    """어떤 파일이든 읽어 공통 구조(dict)로 돌려준다. via 로 방법을 지정할 수 있다 (READERS 의 이름)."""
    p = Path(os.path.expanduser(str(path)))
    if not p.exists():
        return {"error": "missing", "message": f"파일이 없습니다: {p}"}
    ext = p.suffix.lower()
    look = sniff(p)
    routes = [via] if via else guess_routes(ext, look)
    if look == "drm":
        routes = [r for r in routes if not READERS.get(r, (None, True))[1]]  # DRM 은 python 으로 못 연다
    tried = []
    for r in routes:
        if r not in READERS:
            tried.append({"via": r, "error": "unknown_via", "message": f"모르는 방법입니다. 가능한 값: {', '.join(READERS)}"})
            continue
        try:
            d = READERS[r][0](p)
        except Exception as e:  # 한 방법이 실패해도 다음 방법으로 간다
            d = {"error": "reader_failed", "message": f"{type(e).__name__}: {str(e)[:200]}"}
        if not d.get("error") and _garbled(d):
            d = {"error": "garbled", "message": f"{APP_NAME.get(r, r)} 로 열렸지만 글자가 깨져 나왔습니다 (이 형식을 여는 앱이 아닐 수 있음)"}
        if d.get("error"):
            tried.append({"via": r, "error": d["error"], "message": d.get("message")})
            continue
        if d.get("format") == "pptx":
            d.setdefault("via", "python" if r == "pptx" else "powerpoint")
            d["mb"] = round(p.stat().st_size / 1e6, 1)
        d.update({"drm": look == "drm", "path": str(p), "via": d.get("via") or r, "looks_like": look})
        if tried:
            d["tried"] = tried
        return d
    return {"error": "unreadable", "drm": look == "drm", "path": str(p), "looks_like": look,
            "message": f"{p.name} 을(를) 읽지 못했습니다 ({len(tried)}가지 방법 시도).",
            "tried": tried, "hints": _hints(p, look, tried)}


def slide_text(slide):
    return " ".join(sh["text"] for sh in slide["shapes"] if sh.get("text"))

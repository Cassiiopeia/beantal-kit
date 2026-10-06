"""bean-proposal CLI. 모든 서브커맨드는 JSON 한 개를 stdout 에 낸다: ok · code · summary · (데이터) · next.

사용 예 (이 파일이 있는 폴더 기준):
  python proposal_cli.py doctor
  python proposal_cli.py setup-scan
  python proposal_cli.py setup-infer "~/Desktop/제안서"
  python proposal_cli.py setup-save --json '{"project_root": "~/Desktop/제안서"}'
  python proposal_cli.py init --client 고객사 --title 건명 [--due 2026-08-28]
  python proposal_cli.py status <프로젝트 폴더>
  python proposal_cli.py read <파일> [--full]
  python proposal_cli.py profile <예시.pptx>
  python proposal_cli.py outline-check <outline.md> [--requirements requirements.md]
  python proposal_cli.py render <outline.md> --out <결과.pptx> [--project <폴더>]
  python proposal_cli.py ref-add --client 고객사 --title 건명 --year 2026 --file 최종=<pptx>
  python proposal_cli.py ref-migrate
  python proposal_cli.py ref-digest [<id> ...]
  python proposal_cli.py version-new <프로젝트 폴더> --change "수정 이유"
  python proposal_cli.py edit <vNN.pptx> --json '[{"slide":5,"find":"…","replace":"…"}]'
  python proposal_cli.py review <덱.pptx> [--requirements r.md] [--client 고객사] [--md 결과.md]
"""
import argparse
import json
import os
import platform
import re
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


def emit(obj):
    print(json.dumps(obj, ensure_ascii=False, indent=1, default=str))
    return 0 if obj.get("ok") else 1


def expand(p):
    return Path(os.path.expanduser(str(p)))


def need_libs():
    missing = []
    for mod, pkg in (("pptx", "python-pptx"), ("docx", "python-docx"), ("openpyxl", "openpyxl"), ("pypdf", "pypdf")):
        try:
            __import__(mod)
        except ImportError:
            missing.append(pkg)
    return missing


def cmd_doctor(a):
    missing = need_libs()
    ppt = None
    if platform.system() == "Windows":
        ppt = shutil.which("powershell") is not None
    from proposal_lib.setup import CONFIG_PATH, SHARED, load_config
    font_report = None
    if not missing:
        from proposal_lib import fonts as fc
        cfg, _ = load_config()
        wanted = {k: v for k, v in (cfg.get("fonts") or {}).items() if v} or {"title": "G마켓 산스 Bold", "body": "Pretendard Light"}
        _, swaps = fc.resolve(wanted)
        font_report = {"house_fonts": wanted, "missing": swaps,
                       "note": "없는 글꼴은 대체 글꼴로 만듭니다. 받는 사람 PC 에도 없으면 모양이 바뀌므로 PDF 동봉 또는 글꼴 포함 저장을 권장" if swaps else None}
    return emit({
        "fonts": font_report,
        "ok": not missing, "code": "ok" if not missing else "missing_libs",
        "summary": "준비 완료" if not missing else f"필요한 라이브러리가 없습니다: {', '.join(missing)}",
        "python": sys.version.split()[0], "os": platform.system(), "missing": missing,
        "install": f"{Path(sys.executable).name} -m pip install --user {' '.join(missing)}" if missing else None,
        "drm_reader": "PowerPoint (Windows)" if ppt else "없음 — DRM 파일은 읽을 수 없음",
        "config": {"path": str(CONFIG_PATH).replace(str(Path.home()), "~"), "exists": CONFIG_PATH.exists()},
        "company_profile": (SHARED / "company" / "company.json").exists(),
        "next": "setup-scan" if not CONFIG_PATH.exists() else None,
    })


def cmd_setup_scan(a):
    from proposal_lib.setup import scan_roots
    c = scan_roots(a.extra)
    return emit({"ok": True, "code": "scanned", "summary": f"제안서 문서가 모인 폴더 후보 {len(c)}곳", "candidates": c,
                 "next": "후보를 보여 주고 루트를 고르게 한다 (새 폴더 만들기 포함) → setup-infer <루트>"})


def cmd_setup_infer(a):
    from proposal_lib.setup import infer_layout
    r = infer_layout(a.root)
    return emit({"ok": True, "code": "inferred" if r.get("exists") else "new_root", **r,
                 "summary": f"추정 정리 규칙: {r.get('layout_guess')}" if r.get("exists") else "아직 없는 폴더입니다 — 새로 만들 수 있습니다",
                 "next": "추정 규칙이 맞는지 사용자에게 확인 → setup-save"})


def cmd_setup_save(a):
    from proposal_lib.setup import save_config
    values = json.loads(a.json)
    if values.get("project_root"):
        root = expand(values["project_root"])
        if a.create_root:
            root.mkdir(parents=True, exist_ok=True)
    path = save_config(values)
    return emit({"ok": True, "code": "saved", "summary": f"설정을 저장했습니다 ({', '.join(values)})", "path": path.replace(str(Path.home()), "~")})


def cmd_init(a):
    from proposal_lib.project import init_project
    from proposal_lib.setup import load_config
    cfg, exists = load_config()
    root = a.root or cfg.get("project_root")
    if not root:
        return emit({"ok": False, "code": "no_root", "summary": "작업 루트 폴더가 정해지지 않았습니다", "next": "setup-scan"})
    r = init_project(expand(root), a.client, a.title, a.due, layout=cfg.get("project_layout") or "{client}/{yymmdd}_{title}")
    r["next"] = "RFP·자료를 00_입력/ 에 두고 read 로 읽어 requirements.md 작성" if r["ok"] else None
    return emit(r)


def cmd_status(a):
    from proposal_lib.project import status
    return emit(status(expand(a.project)))


def _brief_pptx(d):
    from proposal_lib.review import header_of
    out = []
    for s in d["slides"]:
        h = header_of(s)
        texts = [sh["text"] for sh in s["shapes"] if sh.get("text")]
        out.append({"n": s["n"], **{k: v for k, v in h.items() if v is not None},
                    "kinds": dict(__import__("collections").Counter(sh["kind"] for sh in s["shapes"])),
                    "text": " / ".join(texts)[:600], "notes": s.get("notes", "")[:200]})
    return out


def _lines(d):
    """사람·에이전트가 읽기 쉬운 줄 목록. 제목 수준은 '#' 로, 표는 '| … |' 로 나타낸다."""
    out = []
    if d["format"] == "docx":
        for p in d["paras"]:
            m = re.match(r"(?:Heading|제목)\s*(\d)", p["style"])
            prefix = "#" * int(m.group(1)) + " " if m else ("- " if "List" in p["style"] else "")
            out.append(prefix + p["text"])
        for i, t in enumerate(d["tables"], 1):
            out.append(f"[표 {i}]")
            out += ["| " + " | ".join(r) + " |" for r in t]
    elif d["format"] == "xlsx":
        for name, rows in d["sheets"].items():
            out.append(f"[시트 {name}]")
            out += ["| " + " | ".join(r) + " |" for r in rows]
    elif d["format"] == "pdf":
        for pg in d["pages"]:
            out.append(f"[p{pg['n']}]")
            out += [x for x in pg["text"].splitlines() if x.strip()]
    elif d["format"] == "pptx":
        for s in _brief_pptx(d):
            out.append(f"[슬라이드 {s['n']}] {s.get('num', '')} {s.get('label', '') or ''} {s.get('title', '') or ''}".strip())
            out.append(s["text"])
    return out


def cmd_read(a):
    from proposal_lib.deckio import read_any
    d = read_any(expand(a.file))
    if d.get("error"):
        return emit({"ok": False, "code": d["error"], "summary": d.get("message"), "drm": d.get("drm")})
    if a.lines:
        body = {"lines": _lines(d)}
    elif d["format"] == "pptx" and not a.full:
        body = {"slides": _brief_pptx(d), "size_cm": d["size_cm"], "masters": d.get("masters")}
    else:
        body = {k: v for k, v in d.items() if k not in ("path",)}
    return emit({"ok": True, "code": "read", "summary": f"{d['format']} 읽음 (DRM {'예' if d.get('drm') else '아니오'})", "drm": d.get("drm"), **body})


def cmd_profile(a):
    from proposal_lib.deckio import read_pptx
    from proposal_lib.profile import profile
    d = read_pptx(expand(a.file))
    if d.get("error"):
        return emit({"ok": False, "code": d["error"], "summary": d.get("message")})
    return emit(profile(d))


def cmd_outline_check(a):
    from proposal_lib import outline as ol, tone
    parsed = ol.parse_outline(ol.load(expand(a.outline)))
    reqs = ol.parse_requirements(ol.load(expand(a.requirements))) if a.requirements else None
    findings, coverage = ol.check_outline(parsed, reqs, tone.check_message)
    errors = sum(1 for f in findings if f["sev"] == "error")
    return emit({"ok": errors == 0, "code": "checked", "summary": f"슬라이드 {len(parsed['slides'])}장 · 오류 {errors} · 경고 {sum(1 for f in findings if f['sev'] == 'warn')}",
                 "findings": findings, "coverage": coverage, "next": "오류를 고친 뒤 render" if errors else "render"})


def cmd_snapshot(a):
    import subprocess
    import tempfile
    from proposal_lib.deckio import file_kind
    if platform.system() != "Windows":
        return emit({"ok": False, "code": "unsupported_os", "summary": "스냅샷은 Windows PowerPoint 에서만 지원합니다"})
    src = expand(a.file)
    drm = file_kind(src) == "drm"
    out = expand(a.out) if a.out else Path(tempfile.gettempdir()) / "bean-proposal-snapshots" / src.stem
    if drm and not str(out).startswith(tempfile.gettempdir()):
        return emit({"ok": False, "code": "drm_snapshot_outside_temp", "summary": "DRM 덱의 이미지는 시스템 임시 폴더에만 만들 수 있습니다 (--out 생략)"})
    script = Path(__file__).resolve().parent / "snapshot_pptx.ps1"
    cmd = ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script), "-Path", str(src), "-OutDir", str(out)]
    if a.slides:
        cmd += ["-Slides", a.slides]
    r = subprocess.run(cmd, capture_output=True, timeout=600)
    raw = r.stdout.decode("utf-8", errors="replace")
    try:
        files = json.loads(raw[raw.index("{"):])["files"]
    except (ValueError, KeyError):
        return emit({"ok": False, "code": "snapshot_failed", "summary": (r.stderr.decode("utf-8", errors="replace") or raw)[-300:]})
    if isinstance(files, str):
        files = [files]
    return emit({"ok": True, "code": "snapshot", "summary": f"이미지 {len(files)}장", "dir": str(out), "files": files, "drm": drm,
                 "next": "Read 로 이미지를 보고 레이아웃을 검토" + (" — 검토 후 이 폴더를 지운다 (DRM 덱)" if drm else "")})


def cmd_outline_renumber(a):
    path = expand(a.outline)
    text = path.read_text(encoding="utf-8")
    n = 0
    changes = []

    def repl(m):
        nonlocal n
        n += 1
        if int(m.group(1)) != n:
            changes.append(f"{m.group(1)}→{n:02d}")
        return f"## {n:02d} {m.group(2)}"
    new = re.sub(r"^##\s+(\d{1,3})\s+([·\-|]\s*[A-Z]\d{2}.*)$", repl, text, flags=re.M)
    path.write_text(new, encoding="utf-8")
    return emit({"ok": True, "code": "renumbered", "summary": f"{n}장 · 번호 변경 {len(changes)}건", "changes": changes, "next": "outline-check"})


def cmd_render(a):
    from proposal_lib import outline as ol
    from proposal_lib.render import render
    from proposal_lib.setup import load_config
    cfg, _ = load_config()
    parsed = ol.parse_outline(ol.load(expand(a.outline)))
    out = expand(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    r = render(parsed, cfg, str(out), client_mark=a.client_mark)
    if a.project:
        from proposal_lib.project import record_version
        record_version(expand(a.project), out.relative_to(expand(a.project)) if str(out).startswith(str(expand(a.project))) else out, a.change or "render")
    return emit({"ok": True, "code": "rendered", "summary": f"{r['slides']}장 생성: {out.name}", "path": str(out), **r,
                 "next": f"review \"{out}\""})


def cmd_review(a):
    from proposal_lib.deckio import read_pptx
    from proposal_lib import review as rv, outline as ol
    from proposal_lib.setup import load_config, SHARED
    cfg, _ = load_config()
    d = read_pptx(expand(a.file))
    if d.get("error"):
        return emit({"ok": False, "code": d["error"], "summary": d.get("message")})
    company = None
    cp = SHARED / "company" / "company.json"
    if cp.exists():
        company = json.loads(cp.read_text(encoding="utf-8"))
    clients = None
    refs = Path.home() / ".beantal-kit" / "configs" / "bean-proposal" / "references.json"
    if refs.exists():
        clients = sorted({p["client"] for p in json.loads(refs.read_text(encoding="utf-8")).get("projects", [])})
    reqs = ol.parse_requirements(ol.load(expand(a.requirements))) if a.requirements else None
    r = rv.review(d, cfg, company, clients, a.client, reqs)
    r["drm"] = d.get("drm")
    if a.md:
        md = expand(a.md)
        md.parent.mkdir(parents=True, exist_ok=True)
        md.write_text(rv.to_markdown(r, a.file), encoding="utf-8")
        r["report"] = str(md)
    return emit(r)


def cmd_version_new(a):
    from proposal_lib.project import new_version
    r = new_version(expand(a.project), a.change, a.base)
    if r["ok"]:
        r["next"] = "edit \"%s\" --json '[{\"slide\": 5, \"find\": \"…\", \"replace\": \"…\"}]' 로 이 복사본만 고친다" % r["path"]
    return emit(r)


def cmd_edit(a):
    """복사본의 글자를 바꾼다. 이전 버전 파일에는 쓰지 않는다 (vNN 이 최신이 아니면 거절)."""
    import subprocess, tempfile
    from proposal_lib.project import latest_deck
    f = expand(a.file)
    edits = json.loads(a.json)
    last = latest_deck(f.parent.parent) if f.parent.name == "04_제작" else None
    if last and last.resolve() != f.resolve():
        return emit({"ok": False, "code": "not_latest", "summary": f"최신 버전이 아닙니다 ({last.name}). version-new 로 새 버전을 만든 뒤 고치세요."})
    if platform.system() == "Windows":
        ps1 = Path(__file__).resolve().parent / "edit_pptx.ps1"
        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as t:
            json.dump(edits, t, ensure_ascii=False)
        try:
            r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(ps1), "-Path", str(f), "-EditsJson", t.name],
                               capture_output=True, text=True, encoding="utf-8")
        finally:
            os.unlink(t.name)
        try:
            out = json.loads(r.stdout.strip().splitlines()[-1])
        except Exception:
            return emit({"ok": False, "code": "edit_failed", "summary": (r.stderr or r.stdout)[:400]})
        return emit({"ok": True, "code": "edited", "summary": f"{f.name}: " + ", ".join(f"{x['find'][:12]}→{x['replaced']}곳" for x in out["results"]), **out,
                     "next": f"snapshot \"{f}\" --slides … 로 화면 확인, review 로 점검"})
    from pptx import Presentation
    prs = Presentation(str(f))
    res = []
    for e in edits:
        n = 0
        for i, s in enumerate(prs.slides, 1):
            if e["slide"] not in (0, i):
                continue
            for sh in s.shapes:
                if sh.has_text_frame:
                    for para in sh.text_frame.paragraphs:
                        for run in para.runs:
                            if e["find"] in run.text:
                                run.text = run.text.replace(e["find"], e["replace"]); n += 1
        res.append({"slide": e["slide"], "find": e["find"], "replaced": n})
    prs.save(str(f))
    return emit({"ok": True, "code": "edited", "summary": f"{f.name} 수정", "results": res})


def cmd_ref_add(a):
    from proposal_lib import refs
    files = []
    for spec in a.file:
        role, _, path = spec.partition("=")
        if not path:
            return emit({"ok": False, "code": "bad_file_spec", "summary": f"--file 은 역할=경로 형식입니다: {spec}"})
        files.append((role, path))
    pid = a.id or refs.slug(f"{a.client}-{a.year or ''}")
    proj, log = refs.add(pid, a.client, a.title, a.year, a.quality, files, a.note)
    ok = all(x["ok"] for x in log)
    return emit({"ok": ok, "code": "ref_added", "summary": f"{pid}: {sum(x['ok'] for x in log)}/{len(log)}개 등록", "id": pid, "files": log,
                 "next": f"ref-digest {pid} 로 평문 요약(digest.md)을 만든다"})


def cmd_ref_migrate(a):
    from proposal_lib import refs
    log = refs.migrate()
    return emit({"ok": all(x["ok"] for x in log), "code": "ref_migrated", "summary": f"{sum(x['ok'] for x in log)}/{len(log)}개 복사", "files": log})


def cmd_ref_digest(a):
    from proposal_lib import refs
    from proposal_lib.deckio import read_pptx
    from proposal_lib.profile import profile
    ids = a.id or [p["id"] for p in refs.load()["projects"]]
    res = [refs.digest(i, _brief_pptx, profile, read_pptx) for i in ids]
    return emit({"ok": all(r["ok"] for r in res), "code": "ref_digest", "summary": f"{len(res)}건 처리", "results": res})


def main():
    ap = argparse.ArgumentParser(prog="proposal_cli")
    sp = ap.add_subparsers(dest="cmd", required=True)
    sp.add_parser("doctor").set_defaults(fn=cmd_doctor)
    p = sp.add_parser("setup-scan"); p.add_argument("--extra", nargs="*"); p.set_defaults(fn=cmd_setup_scan)
    p = sp.add_parser("setup-infer"); p.add_argument("root"); p.set_defaults(fn=cmd_setup_infer)
    p = sp.add_parser("setup-save"); p.add_argument("--json", required=True); p.add_argument("--create-root", action="store_true"); p.set_defaults(fn=cmd_setup_save)
    p = sp.add_parser("init"); p.add_argument("--client", required=True); p.add_argument("--title", required=True); p.add_argument("--due"); p.add_argument("--root"); p.set_defaults(fn=cmd_init)
    p = sp.add_parser("status"); p.add_argument("project"); p.set_defaults(fn=cmd_status)
    p = sp.add_parser("read"); p.add_argument("file"); p.add_argument("--full", action="store_true"); p.add_argument("--lines", action="store_true"); p.set_defaults(fn=cmd_read)
    p = sp.add_parser("profile"); p.add_argument("file"); p.set_defaults(fn=cmd_profile)
    p = sp.add_parser("outline-check"); p.add_argument("outline"); p.add_argument("--requirements"); p.set_defaults(fn=cmd_outline_check)
    p = sp.add_parser("snapshot"); p.add_argument("file"); p.add_argument("--out"); p.add_argument("--slides"); p.set_defaults(fn=cmd_snapshot)
    p = sp.add_parser("outline-renumber"); p.add_argument("outline"); p.set_defaults(fn=cmd_outline_renumber)
    p = sp.add_parser("render"); p.add_argument("outline"); p.add_argument("--out", required=True); p.add_argument("--project"); p.add_argument("--change"); p.add_argument("--client-mark"); p.set_defaults(fn=cmd_render)
    p = sp.add_parser("ref-add"); p.add_argument("--client", required=True); p.add_argument("--title", required=True); p.add_argument("--year"); p.add_argument("--id"); p.add_argument("--quality", default="reference", choices=["gold", "reference", "wip"]); p.add_argument("--file", action="append", required=True, help="역할=경로 (여러 번)"); p.add_argument("--note"); p.set_defaults(fn=cmd_ref_add)
    sp.add_parser("ref-migrate").set_defaults(fn=cmd_ref_migrate)
    p = sp.add_parser("ref-digest"); p.add_argument("id", nargs="*"); p.set_defaults(fn=cmd_ref_digest)
    p = sp.add_parser("version-new"); p.add_argument("project"); p.add_argument("--change", required=True); p.add_argument("--base"); p.set_defaults(fn=cmd_version_new)
    p = sp.add_parser("edit"); p.add_argument("file"); p.add_argument("--json", required=True, help='[{"slide":5,"find":"…","replace":"…"}]'); p.set_defaults(fn=cmd_edit)
    p = sp.add_parser("review"); p.add_argument("file"); p.add_argument("--requirements"); p.add_argument("--client"); p.add_argument("--md"); p.set_defaults(fn=cmd_review)
    a = ap.parse_args()
    missing = need_libs() if a.cmd != "doctor" else []
    if missing:
        return emit({"ok": False, "code": "missing_libs", "summary": f"필요한 라이브러리가 없습니다: {', '.join(missing)}",
                     "install": f"python -m pip install --user {' '.join(missing)}"})
    try:
        return a.fn(a)
    except Exception as e:  # 예상 못 한 오류도 JSON 으로 돌려준다
        return emit({"ok": False, "code": "exception", "summary": f"{type(e).__name__}: {e}"})


if __name__ == "__main__":
    sys.exit(main())
